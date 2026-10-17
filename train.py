from __future__ import annotations

import argparse
import random
import re
import sys
from pathlib import Path

import cv2
import joblib
import numpy as np
from openpyxl import load_workbook
from sklearn.ensemble import ExtraTreesClassifier
from sklearn.metrics import classification_report
from sklearn.metrics import precision_recall_fscore_support
from sklearn.model_selection import GroupShuffleSplit

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.model import MODEL_PATH, frame_features


def canonical(value: str) -> str:
    return value.strip().lower().replace("х", "x").replace("Х", "x")


def timestamp_seconds(value: object) -> float | None:
    if value is None:
        return None
    text = str(value).lower().strip()
    minute = re.search(r"(\d+)\s*минут", text)
    second = re.search(r"(\d+)\s*секунд", text)
    total = (int(minute.group(1)) * 60 if minute else 0) + (int(second.group(1)) if second else 0)
    return float(total) if minute or second else None


def read_labels(dataset: Path) -> dict[str, tuple[int, float | None]]:
    book = load_workbook(dataset / "open_labels.xlsx", read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    result = {}
    for name, timestamp, label in list(sheet.values)[1:]:
        if not name or not label:
            continue
        is_positive = "нарушение" in str(label).lower() and "без наруш" not in str(label).lower()
        result[canonical(Path(str(name)).stem)] = (int(is_positive), timestamp_seconds(timestamp))
    return result


def extract_video(path: Path, is_positive: int, event_time: float | None, sample_fps: float):
    cap = cv2.VideoCapture(str(path))
    if not cap.isOpened():
        print(f"Пропускаю нечитаемый файл: {path.name}", file=sys.stderr)
        return []
    fps = cap.get(cv2.CAP_PROP_FPS) or 29.0
    stride = max(1, round(fps / sample_fps))
    rows, labels, groups = [], [], []
    previous = None
    frame_number = 0
    while True:
        ok, frame = cap.read()
        if not ok:
            break
        if frame_number % stride == 0:
            t = frame_number / fps
            positive_window = bool(is_positive and event_time is not None and abs(t - event_time) <= 2.0)
            background = not is_positive or event_time is None or abs(t - event_time) >= 6.0
            if positive_window or background:
                rows.append(frame_features(frame, previous))
                labels.append(int(positive_window))
                groups.append(path.name)
            previous = cv2.resize(cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY), (96, 96), interpolation=cv2.INTER_AREA)
        frame_number += 1
    cap.release()
    return rows, labels, groups


def main() -> None:
    parser = argparse.ArgumentParser(description="Обучение локального детектора по меткам open набора")
    parser.add_argument("--dataset", required=True, type=Path, help="Распакованная папка open/ из архива")
    parser.add_argument("--sample-fps", type=float, default=1.0)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--threshold", type=float, default=0.45, help="Минимальный порог события (не может быть ниже 0.45)")
    args = parser.parse_args()
    labels_by_name = read_labels(args.dataset)
    videos_dir = args.dataset / "videos"
    if not videos_dir.exists():
        videos_dir = args.dataset
    videos = {canonical(p.stem): p for p in videos_dir.glob("*.mov")}
    missing = sorted(set(labels_by_name) - set(videos))
    if missing:
        print(f"Не найдены в папке videos: {', '.join(missing)}", file=sys.stderr)
    x, y, groups = [], [], []
    for normalized, path in sorted(videos.items()):
        if normalized not in labels_by_name:
            continue
        positive, event_time = labels_by_name[normalized]
        result = extract_video(path, positive, event_time, args.sample_fps)
        if not result:
            continue
        rows, labels, source_groups = result
        x.extend(rows); y.extend(labels); groups.extend(source_groups)
        print(f"{path.name}: метка={'нарушение' if positive else 'без нарушения'}, окон={len(labels)}")
    if len(set(y)) < 2:
        raise SystemExit("Недостаточно примеров обоих классов. Проверьте архив и open_labels.xlsx.")
    x = np.asarray(x, dtype=np.float32)
    y = np.asarray(y, dtype=np.int8)
    groups = np.asarray(groups)
    print(f"Выборка: {len(y)} окон, положительных {int(y.sum())}, отрицательных {int((y == 0).sum())}")

    splitter = GroupShuffleSplit(n_splits=1, test_size=0.25, random_state=args.seed)
    train_idx, valid_idx = next(splitter.split(x, y, groups))
    validation_model = ExtraTreesClassifier(
        n_estimators=240, max_depth=16, min_samples_leaf=2,
        class_weight="balanced", n_jobs=-1, random_state=args.seed,
    )
    validation_model.fit(x[train_idx], y[train_idx])
    probabilities = validation_model.predict_proba(x[valid_idx])[:, 1]
    y_valid = y[valid_idx]
    print("Групповая проверка по отдельным видео (ориентир, малый набор):")
    threshold = max(0.45, args.threshold)
    precision, recall, f2, _ = precision_recall_fscore_support(
        y_valid, probabilities >= threshold, beta=2, average="binary", zero_division=0
    )
    print(f"Качество на пороге {threshold:.0%}: precision={precision:.3f}, recall={recall:.3f}, F2={f2:.3f}")
    print(classification_report(y_valid, probabilities >= threshold, zero_division=0))

    final_model = ExtraTreesClassifier(
        n_estimators=320, max_depth=16, min_samples_leaf=2,
        class_weight="balanced", n_jobs=-1, random_state=args.seed,
    )
    final_model.fit(x, y)
    MODEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump({
        "classifier": final_model, "sample_fps": args.sample_fps,
        "feature_version": 1, "threshold": threshold,
        "validation": {"precision": float(precision), "recall": float(recall), "f2": float(f2), "threshold": threshold, "video_groups": int(len(np.unique(groups[valid_idx]))), "positive_windows": int(y_valid.sum()), "negative_windows": int((y_valid == 0).sum())},
    }, MODEL_PATH, compress=3)
    print(f"Модель сохранена: {MODEL_PATH}")
    print("Порог не может быть ниже 45%. Подтвердите качество на отдельном наборе перед демонстрацией.")


if __name__ == "__main__":
    main()

